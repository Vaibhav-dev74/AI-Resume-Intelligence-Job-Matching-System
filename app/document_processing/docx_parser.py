import io
from typing import Tuple


class DOCXParser:
    @classmethod
    def extract_text(cls, file_bytes: bytes) -> Tuple[str, int]:
        try:
            import docx
        except ImportError:
            raise RuntimeError("python-docx is not installed. Please install 'python-docx'.")

        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            raise ValueError(f"Failed to open DOCX document: {e}")

        paragraphs = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if text:
                paragraphs.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    deduped = []
                    for c in row_cells:
                        if not deduped or deduped[-1] != c:
                            deduped.append(c)
                    paragraphs.append(" | ".join(deduped))

        full_text = "\n".join(paragraphs)
        return full_text.strip(), 1


docx_parser = DOCXParser()
